from collections.abc import Callable
from pathlib import Path

import docx
from docx.document import Document as DocxDocument

from app.adapters.parsing.docx import _extract_section_code, parse_docx


def _make_docx(tmp_path: Path, build: Callable[[DocxDocument], object]) -> Path:
    document = docx.Document()
    build(document)
    path = tmp_path / "test.docx"
    document.save(str(path))
    return path


def test_extracts_paragraph_text_with_page_one_and_no_bbox(tmp_path: Path) -> None:
    path = _make_docx(tmp_path, lambda d: d.add_paragraph("Plain body text"))

    parsed = parse_docx(path)

    assert len(parsed.blocks) == 1
    block = parsed.blocks[0]
    assert block.text == "Plain body text"
    assert block.page == 1
    assert block.kind == "text"
    assert block.is_title is False


def test_detects_heading_by_style_not_by_font_size(tmp_path: Path) -> None:
    def build(d: DocxDocument) -> None:
        d.add_heading("3.2 Lenkkraft", level=2)
        d.add_paragraph("Normal body sentence.")

    parsed = parse_docx(_make_docx(tmp_path, build))

    titles = {b.text: b.is_title for b in parsed.blocks}
    assert titles["3.2 Lenkkraft"] is True
    assert titles["Normal body sentence."] is False


def test_numbered_heading_section_code_strips_to_the_number(tmp_path: Path) -> None:
    path = _make_docx(tmp_path, lambda d: d.add_heading("3.2 Lenkkraft", level=2))

    parsed = parse_docx(path)

    assert parsed.blocks[0].section_path == "3.2"


def test_non_numbered_heading_uses_its_own_text_as_the_code() -> None:
    assert _extract_section_code("Anhang A: Änderungshistorie") == "Anhang A: Änderungshistorie"


def test_section_path_is_inherited_by_following_paragraphs(tmp_path: Path) -> None:
    def build(d: DocxDocument) -> None:
        d.add_heading("1. Einleitung", level=1)
        d.add_paragraph("Body under section 1.")
        d.add_heading("1.1 Zweck", level=2)
        d.add_paragraph("Body under section 1.1.")

    parsed = parse_docx(_make_docx(tmp_path, build))

    by_text = {b.text: b.section_path for b in parsed.blocks}
    assert by_text["1. Einleitung"] == "1"
    assert by_text["Body under section 1."] == "1"
    assert by_text["1.1 Zweck"] == "1.1"
    assert by_text["Body under section 1.1."] == "1.1"


def test_bold_run_detected_on_a_non_heading_paragraph(tmp_path: Path) -> None:
    def build(d: DocxDocument) -> None:
        p = d.add_paragraph()
        run = p.add_run("LH-3.2.1 [MUSS] ")
        run.bold = True
        p.add_run("Rest of the requirement text.")

    parsed = parse_docx(_make_docx(tmp_path, build))

    assert parsed.blocks[0].bold is True


def test_table_extracted_as_a_markdown_block_after_the_preceding_heading(
    tmp_path: Path,
) -> None:
    def build(d: DocxDocument) -> None:
        d.add_heading("5. Umgebungsbedingungen", level=1)
        table = d.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Parameter"
        table.rows[0].cells[1].text = "Wert"
        row = table.add_row().cells
        row[0].text = "Betriebsspannung"
        row[1].text = "9 bis 16 V"

    parsed = parse_docx(_make_docx(tmp_path, build))

    table_blocks = [b for b in parsed.blocks if b.kind == "table"]
    assert len(table_blocks) == 1
    lines = table_blocks[0].text.splitlines()
    assert lines[0] == "| Parameter | Wert |"
    assert lines[1] == "| --- | --- |"
    assert lines[2] == "| Betriebsspannung | 9 bis 16 V |"
    assert table_blocks[0].section_path == "5"


def test_paragraphs_and_tables_stay_in_document_order(tmp_path: Path) -> None:
    def build(d: DocxDocument) -> None:
        d.add_paragraph("Before the table.")
        table = d.add_table(rows=2, cols=1)
        table.rows[0].cells[0].text = "Header"
        table.rows[1].cells[0].text = "In the table."
        d.add_paragraph("After the table.")

    parsed = parse_docx(_make_docx(tmp_path, build))

    assert [b.text.splitlines()[-1] for b in parsed.blocks] == [
        "Before the table.",
        "| In the table. |",
        "After the table.",
    ]
