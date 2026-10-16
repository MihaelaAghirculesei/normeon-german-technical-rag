"""Upload inspection and the request-size limit (Giorno 26), on real files:
PDFs and Word documents are generated here, including an encrypted PDF and
a zip bomb, and the size limit is exercised over HTTP."""

import io
import zipfile
from collections.abc import Iterator

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient
from starlette.applications import Starlette
from starlette.requests import Request
from starlette.responses import JSONResponse
from starlette.routing import Route

from app.adapters.parsing.inspect import UploadLimits, inspect_upload
from app.core.errors import (
    EncryptedDocumentError,
    MalformedDocumentError,
    SuspiciousArchiveError,
    TooManyPagesError,
)
from app.core.limits import BodySizeLimit
from app.main import app

LIMITS = UploadLimits(max_pages=3, max_uncompressed_bytes=20 * 2**20, max_compression_ratio=100)


def _pdf(pages: int = 1, **save_options: object) -> bytes:
    doc = pymupdf.open()
    for n in range(pages):
        doc.new_page().insert_text((50, 72), f"Seite {n + 1}: Lenkkraft 300 N")
    data: bytes = doc.tobytes(**save_options)
    doc.close()
    return data


def _docx() -> bytes:
    document = docx.Document()
    document.add_heading("1 Allgemeines", level=1)
    document.add_paragraph("Die Lenkkraft darf 250 N nicht ueberschreiten.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _zip(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()


# -- PDF ------------------------------------------------------------------------


def test_a_readable_pdf_passes() -> None:
    inspect_upload(".pdf", _pdf(pages=3), LIMITS)


def test_a_password_protected_pdf_is_refused() -> None:
    encrypted = _pdf(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="geheim", owner_pw="geheim")

    with pytest.raises(EncryptedDocumentError):
        inspect_upload(".pdf", encrypted, LIMITS)


@pytest.mark.parametrize("content", [b"", b"%PDF-1.4 kaputt", b"\x00" * 512])
def test_a_broken_pdf_is_refused(content: bytes) -> None:
    with pytest.raises(MalformedDocumentError):
        inspect_upload(".pdf", content, LIMITS)


def test_a_pdf_over_the_page_limit_is_refused() -> None:
    with pytest.raises(TooManyPagesError, match="4 pages"):
        inspect_upload(".pdf", _pdf(pages=4), LIMITS)


# -- DOCX -----------------------------------------------------------------------


def test_a_real_word_document_passes() -> None:
    inspect_upload(".docx", _docx(), LIMITS)


def test_a_password_protected_word_document_is_refused() -> None:
    # Word wraps an encrypted .docx in an OLE compound file.
    ole = bytes.fromhex("d0cf11e0a1b11ae1") + b"\x00" * 504

    with pytest.raises(EncryptedDocumentError):
        inspect_upload(".docx", ole, LIMITS)


def test_a_file_that_is_no_zip_is_refused() -> None:
    with pytest.raises(MalformedDocumentError):
        inspect_upload(".docx", b"das ist kein Word-Dokument", LIMITS)


def test_a_zip_without_a_document_body_is_refused() -> None:
    with pytest.raises(MalformedDocumentError, match="no Word document body"):
        inspect_upload(".docx", _zip({"notes.txt": b"hallo"}), LIMITS)


def test_a_zip_bomb_is_refused_without_unpacking_it() -> None:
    # 16 MB of zeros deflate to a few KB: a ratio a document never has.
    bomb = _zip({"word/document.xml": b"<w/>", "word/media/x.bin": b"\x00" * 16 * 2**20})
    assert len(bomb) < 100 * 2**10

    with pytest.raises(SuspiciousArchiveError, match="zip bomb"):
        inspect_upload(".docx", bomb, LIMITS)


def test_a_word_file_unpacking_past_the_size_limit_is_refused() -> None:
    small = UploadLimits(max_pages=3, max_uncompressed_bytes=2**20, max_compression_ratio=10**6)
    big = _zip({"word/document.xml": b"<w/>" * (2**20 // 2)})

    with pytest.raises(SuspiciousArchiveError, match="unpack"):
        inspect_upload(".docx", big, small)


# -- request size ---------------------------------------------------------------


@pytest.fixture
def limited() -> Iterator[TestClient]:
    async def echo(request: Request) -> JSONResponse:
        return JSONResponse({"received": len(await request.body())})

    inner = Starlette(
        routes=[
            Route("/api/v1/documents", echo, methods=["POST"]),
            Route("/api/v1/chat", echo, methods=["POST"]),
        ]
    )
    with TestClient(BodySizeLimit(inner, max_upload_bytes=1000, max_request_bytes=100)) as client:
        yield client


def test_a_body_within_the_limit_reaches_the_app(limited: TestClient) -> None:
    response = limited.post("/api/v1/documents", content=b"x" * 1000)

    assert response.status_code == 200
    assert response.json() == {"received": 1000}


def test_an_upload_over_the_limit_is_refused_from_its_headers(limited: TestClient) -> None:
    response = limited.post("/api/v1/documents", content=b"x" * 1001)

    assert response.status_code == 413
    assert response.json()["code"] == "upload_too_large"


def test_other_endpoints_have_a_smaller_limit(limited: TestClient) -> None:
    assert limited.post("/api/v1/chat", content=b"x" * 100).status_code == 200
    assert limited.post("/api/v1/chat", content=b"x" * 101).status_code == 413


def test_a_body_without_a_declared_length_is_refused(limited: TestClient) -> None:
    def chunks() -> Iterator[bytes]:
        yield b"x" * 10

    response = limited.post("/api/v1/documents", content=chunks())

    assert response.status_code == 411
    assert response.json()["code"] == "length_required"


def test_the_app_refuses_a_200_mb_upload_before_reading_it() -> None:
    # Only the header claims 200 MB; nothing that large is ever sent. The
    # refusal comes before authentication, since nothing is read.
    response = TestClient(app).post(
        "/api/v1/documents",
        content=b"x",
        headers={"content-length": str(200 * 2**20), "content-type": "multipart/form-data"},
    )

    assert response.status_code == 413
