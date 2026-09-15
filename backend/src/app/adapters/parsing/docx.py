import re
from collections.abc import Iterator
from dataclasses import replace
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.domain.parsing import Block, BlockKind, ParsedDocument

# A docx has no page concept the way a rendered PDF does -- every block is
# stamped with the same page, matching the whole-document span a caller
# would show if it cited this source. Fine for the short Lastenheft-sized
# documents this parser targets; revisit only if a much longer docx needs
# page-level citation granularity.
_DOCX_PAGE = 1

_LEADING_NUMBER_RE = re.compile(r"^(\d+(?:\.\d+)*)")
_HEADING_LEVEL_RE = re.compile(r"^(Title|Heading \d+)$")


def parse_docx(path: str | Path) -> ParsedDocument:
    document = docx.Document(str(path))
    blocks: list[Block] = []
    for item in _iter_block_items(document):
        if isinstance(item, Table):
            markdown = _table_to_markdown(item)
            if markdown:
                blocks.append(_block(markdown, kind="table"))
            continue

        text = item.text.strip()
        if not text:
            continue
        style_name = item.style.name if item.style is not None else None
        is_title = bool(_HEADING_LEVEL_RE.match(style_name or ""))
        bold = not is_title and any(run.bold for run in item.runs)
        blocks.append(_block(text, is_title=is_title, bold=bold))

    return ParsedDocument(blocks=_assign_section_paths(blocks))


def _block(
    text: str, *, kind: BlockKind = "text", is_title: bool = False, bold: bool = False
) -> Block:
    return Block(
        text=text,
        page=_DOCX_PAGE,
        bbox=(0.0, 0.0, 0.0, 0.0),
        font_size=0.0,
        bold=bold,
        kind=kind,
        is_title=is_title,
    )


def _iter_block_items(document: DocxDocument) -> Iterator[Paragraph | Table]:
    """Paragraphs and tables in document order -- python-docx groups them
    into separate collections (`document.paragraphs` / `.tables`) that lose
    the interleaving, so this walks the underlying XML body directly."""
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _table_to_markdown(table: Table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""

    header, *body = rows
    col_count = len(header)
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * col_count) + " |",
    ]
    for row in body:
        padded = row + [""] * (col_count - len(row))
        lines.append("| " + " | ".join(padded[:col_count]) + " |")
    return "\n".join(lines)


def _extract_section_code(text: str) -> str:
    """A leading `\\d+(.\\d+)*` heading number ("3.2 Lenkkraft" -> "3.2"), or
    the heading's own text when there isn't one ("Anhang A: ..." stays its
    own top-level path -- there's nothing to nest it under)."""
    match = _LEADING_NUMBER_RE.match(text)
    return match.group(1) if match else text


def _assign_section_paths(blocks: list[Block]) -> list[Block]:
    current_path: str | None = None
    result: list[Block] = []
    for b in blocks:
        if b.is_title:
            current_path = _extract_section_code(b.text)
        result.append(replace(b, section_path=current_path))
    return result
