"""Generates the two synthetic Lastenheft versions used as corpus documents.

Not part of the backend application — a one-off authoring tool so v1.2 and
v2.0 stay derived from a single source instead of hand-edited in parallel
and drifting apart by accident. Needs `pip install python-docx` to run.

Usage: python generate_lastenheft.py
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


@dataclass(frozen=True)
class Variant:
    version_label: str
    valid_from: str
    valid_until: str | None
    max_lenkkraft_n: int
    diagnose_zeit_ms: int
    changelog: str


V1_2 = Variant(
    version_label="v1.2",
    valid_from="2024-01-01",
    valid_until="2025-12-31",
    max_lenkkraft_n=300,
    diagnose_zeit_ms=150,
    changelog=(
        "v1.2 (2024-01-01): Anpassung der Diagnose-Ansprechzeit nach "
        "Rückmeldung aus der Erprobung (LH-6.2.1). Redaktionelle "
        "Korrekturen in Abschnitt 5."
    ),
)

V2_0 = Variant(
    version_label="v2.0",
    valid_from="2026-01-01",
    valid_until=None,
    max_lenkkraft_n=250,
    diagnose_zeit_ms=100,
    changelog=(
        "v2.0 (2026-01-01): Verschärfung der maximal zulässigen Lenkkraft "
        "bei Ausfall der Servounterstützung von 300 N auf 250 N (LH-3.2.1), "
        "auf Basis der aktualisierten Grenzwertbetrachtung zu UN R79. "
        "Verkürzung der maximalen Diagnose-Ansprechzeit von 150 ms auf "
        "100 ms (LH-6.2.1). Diese Version ersetzt v1.2 vollständig; "
        "abweichende Werte in v1.2 gelten ab dem 2026-01-01 nicht mehr."
    ),
)


def add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def add_req(doc: Document, req_id: str, text: str, modal: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"{req_id} [{modal}] ")
    run.bold = True
    p.add_run(text)


def build(variant: Variant, out_path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    title = doc.add_heading("Lastenheft — Elektrisches Servolenksystem (EPS)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("für Fahrzeuge der Kompaktklasse")

    meta = doc.add_table(rows=0, cols=2)
    for label, value in [
        ("Version", variant.version_label),
        ("Gültig ab", variant.valid_from),
        ("Gültig bis", variant.valid_until or "bis auf Widerruf"),
        ("Dokumenttyp", "Lastenheft (synthetisch, zu Demonstrationszwecken)"),
        ("Sprache", "Deutsch"),
        ("Bezug", "UN R79 (Lenkanlagen), StVZO"),
    ]:
        row = meta.add_row().cells
        row[0].text = label
        row[1].text = value

    add_heading(doc, "1. Einleitung und Geltungsbereich", 1)
    add_heading(doc, "1.1 Zweck des Dokuments", 2)
    doc.add_paragraph(
        "Dieses Lastenheft beschreibt die funktionalen, sicherheitsbezogenen "
        "und umgebungsbezogenen Anforderungen an das elektrische "
        "Servolenksystem (Electric Power Steering, EPS) für Fahrzeuge der "
        "Kompaktklasse. Es dient als verbindliche Grundlage für Entwicklung, "
        "Erprobung und Abnahme des Systems durch den Lieferanten."
    )
    add_heading(doc, "1.2 Anwendungsbereich", 2)
    doc.add_paragraph(
        "Das Lastenheft gilt für alle EPS-Varianten der Baureihe, die für "
        "den Einsatz in Mitgliedstaaten der Europäischen Union unter "
        "Anwendung der UN-Regelung Nr. 79 typgenehmigt werden sollen."
    )

    add_heading(doc, "1.3 Begriffe und Abkürzungen", 2)
    doc.add_paragraph(
        "MUSS kennzeichnet eine zwingende Anforderung. SOLL kennzeichnet "
        "eine dringend empfohlene Anforderung, von der nur mit begründeter "
        "Ausnahme abgewichen werden darf. KANN kennzeichnet eine optionale "
        "Anforderung."
    )

    add_heading(doc, "2. Referenzierte Dokumente", 1)
    doc.add_paragraph(
        "UN R79 — Einheitliche Bedingungen für die Genehmigung der "
        "Fahrzeuge hinsichtlich der Lenkanlage.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "StVZO — Straßenverkehrs-Zulassungs-Ordnung, insbesondere die "
        "zulassungsrelevanten Vorschriften zu Lenkeinrichtungen.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "UN R155 — Cybersicherheit und Cybersicherheits-Managementsystem.",
        style="List Bullet",
    )

    add_heading(doc, "3. Funktionale Anforderungen an das Lenksystem", 1)
    add_heading(doc, "3.1 Lenkübersetzung", 2)
    add_req(
        doc,
        "LH-3.1.1",
        "Die Lenkübersetzung muss über den gesamten Geschwindigkeitsbereich "
        "so ausgelegt sein, dass maximal 2,5 Lenkradumdrehungen von "
        "Anschlag zu Anschlag erforderlich sind.",
        "MUSS",
    )

    add_heading(doc, "3.2 Lenkkraft", 2)
    add_req(
        doc,
        "LH-3.2.1",
        f"Bei Ausfall der Servounterstützung darf die vom Fahrer an der "
        f"Lenkradfelge aufzubringende Lenkkraft einen Wert von "
        f"{variant.max_lenkkraft_n} N nicht überschreiten, um das Fahrzeug "
        f"mit einer Verzögerung von 2 m/s² aus einer Kurvenfahrt mit 0,4 g "
        f"Querbeschleunigung sicher zum Stillstand zu bringen "
        f"(vgl. UN R79, Abschnitt 5.1.2).",
        "MUSS",
    )
    add_req(
        doc,
        "LH-3.2.2",
        "Die Lenkkraft im Normalbetrieb (Servounterstützung aktiv) muss "
        "geschwindigkeitsabhängig moduliert werden, um bei Parkiervorgängen "
        "ein niedriges und bei Autobahnfahrt ein höheres Rückstellmoment "
        "zu erzeugen.",
        "MUSS",
    )

    add_heading(doc, "3.3 Rückstellverhalten", 2)
    add_req(
        doc,
        "LH-3.3.1",
        "Das Lenksystem soll nach einer Kurvenfahrt ohne Fahrereingriff in "
        "die Geradeausstellung zurückkehren.",
        "SOLL",
    )

    add_heading(doc, "4. Sicherheitsanforderungen", 1)
    add_req(
        doc,
        "LH-4.1.1",
        "Das EPS muss im Fehlerfall in einen sicheren Zustand mit "
        "manueller Lenkbarkeit übergehen; ein vollständiger Verlust der "
        "Lenkbarkeit ist unzulässig.",
        "MUSS",
    )
    add_req(
        doc,
        "LH-4.2.1",
        "Das Cybersicherheits-Managementsystem für die Softwarekomponenten "
        "des EPS muss die Anforderungen der UN R155 erfüllen.",
        "MUSS",
    )

    add_heading(doc, "5. Umgebungsbedingungen", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Parameter"
    hdr[1].text = "Wert"
    hdr[2].text = "Einheit"
    hdr[3].text = "Anforderungs-ID"

    rows = [
        ("Betriebstemperaturbereich", "-40 bis +85", "°C", "LH-5.1.1"),
        (
            "Maximal zulässige Lenkkraft bei Ausfall",
            str(variant.max_lenkkraft_n),
            "N",
            "LH-3.2.1",
        ),
        (
            "Maximale Diagnose-Ansprechzeit",
            str(variant.diagnose_zeit_ms),
            "ms",
            "LH-6.2.1",
        ),
        ("Betriebsspannungsbereich", "9 bis 16", "V", "LH-5.2.1"),
        ("Schutzart Steuergerät", "IP6K9K", "-", "LH-5.3.1"),
    ]
    for param, value, unit, req_id in rows:
        cells = table.add_row().cells
        cells[0].text = param
        cells[1].text = value
        cells[2].text = unit
        cells[3].text = req_id

    add_heading(doc, "6. Diagnose und Fehlerüberwachung", 1)
    add_req(
        doc,
        "LH-6.1.1",
        "Das Steuergerät muss den Ausfall der Servounterstützung "
        "innerhalb einer Fahrzeugzyklen-Grenze erkennen und dem Fahrer "
        "über eine Warnleuchte anzeigen.",
        "MUSS",
    )
    add_req(
        doc,
        "LH-6.2.1",
        f"Die maximale Ansprechzeit der Diagnosefunktion vom Auftreten des "
        f"Fehlers bis zur Aktivierung der Warnleuchte darf "
        f"{variant.diagnose_zeit_ms} ms nicht überschreiten.",
        "MUSS",
    )

    add_heading(doc, "7. Dokumentation und Nachweisführung", 1)
    add_req(
        doc,
        "LH-7.1.1",
        "Der Lieferant muss für jede Anforderung dieses Lastenhefts einen "
        "Nachweis (Prüfbericht, Simulation oder Berechnung) führen und auf "
        "Anfrage vorlegen.",
        "MUSS",
    )

    add_heading(doc, "Anhang A: Änderungshistorie", 1)
    doc.add_paragraph(variant.changelog)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"wrote {out_path}")


if __name__ == "__main__":
    base = Path(__file__).parent
    build(V1_2, base / "Lastenheft-EPS-v1.2.docx")
    build(V2_0, base / "Lastenheft-EPS-v2.0.docx")
